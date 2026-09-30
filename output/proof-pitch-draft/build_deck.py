"""Build one editable deck, a PDF, and a self-contained HTML pitch from one layout."""
from pathlib import Path
from html import escape
from io import BytesIO
import base64
import json
import shutil
import math

from PIL import Image
from fontTools.ttLib import TTFont
from fontTools.pens.svgPathPen import SVGPathPen
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as RLFont
from reportlab.lib.utils import ImageReader
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
W, H = 1280, 720
INK, MUTED, FOREST, SAGE, PAPER, LINE = '1D3027', '5C6D63', '245C48', 'E7EFE9', 'FAFBF9', 'DCE5DF'
WHITE, LIGHT, AMBER, AMBER_BG = 'FFFFFF', 'D8EBE2', '996A3E', 'F2E8DB'
FONTS = {
    'serif': ('newsreader', '400', 'normal', 'Newsreader'),
    'italic': ('newsreader', '400', 'italic', 'Newsreader'),
    'sans': ('dm-sans', '400', 'normal', 'DM Sans'),
    'medium': ('dm-sans', '500', 'normal', 'DM Sans'),
}
FONT_DATA = {}
for name, (family, weight, style, display) in FONTS.items():
    stem = f'{family}-latin-{weight}-{style}'
    source = ROOT / 'node_modules' / '@fontsource' / family / 'files'
    font = TTFont(source / f'{stem}.woff')
    font.flavor = None
    target = ASSETS / f'{stem}.ttf'
    font.save(target)
    pdfmetrics.registerFont(RLFont(name, str(target)))
    FONT_DATA[name] = base64.b64encode((source / f'{stem}.woff2').read_bytes()).decode()
    license_path = ROOT / 'node_modules' / '@fontsource' / family / 'LICENSE'
    if license_path.exists():
        shutil.copyfile(license_path, ASSETS / f'{family}-LICENSE.txt')

logo = Image.open(ROOT / 'public/images/proof-logo-drawn-v1.png').convert('RGBA')
logo.thumbnail((300, 300), Image.Resampling.LANCZOS)
logo.save(ASSETS / 'proof-logo.png', optimize=True)
shutil.copyfile(ROOT / 'public/design-system/app-reference.jpg', ASSETS / 'proof-interface.jpg')

slides = []
def slide(section, title, bg=PAPER):
    s = {'section': section, 'title': title, 'bg': bg, 'elements': [], 'notes': ''}
    slides.append(s)
    return s
def add(s, type, **kwargs):
    s['elements'].append({'type': type, **kwargs})
def text(s, value, x, y, w, size=24, font='sans', color=INK, h=None, leading=1.18, align='left', url=None):
    count = len(value.split('\n'))
    add(s, 'text', text=value, x=x, y=y, w=w, h=h or size * leading * count + 7,
        size=size, font=font, color=color, leading=leading, align=align, url=url)
def rect(s, x, y, w, h, fill=WHITE, stroke=None, radius=16):
    add(s, 'rect', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, radius=radius)
def line(s, x1, y1, x2, y2, color=LINE, width=2):
    add(s, 'line', x1=x1, y1=y1, x2=x2, y2=y2, color=color, width=width)
def image(s, path, x, y, w, h):
    add(s, 'image', path=str(path), x=x, y=y, w=w, h=h)
def chrome(s):
    dark = s['bg'] == FOREST
    color, mute = (WHITE, LIGHT) if dark else (INK, MUTED)
    text(s, 'proof.', 64, 35, 150, 33, 'serif', color)
    text(s, s['section'], 800, 42, 416, 16, 'sans', mute, align='right')
    line(s, 64, 666, 1216, 666, '487560' if dark else LINE, 1)
    text(s, 'Proof · Hackathon pitch · First draft', 64, 680, 650, 13, 'sans', mute)
    text(s, f'{len(slides):02d} / 08', 1070, 680, 146, 13, 'sans', mute, align='right')
def arrow(s, x1, y1, x2, y2, color=FOREST):
    line(s, x1, y1, x2, y2, color, 2)
    ang = math.atan2(y2-y1, x2-x1)
    for offset in (-.5, .5):
        line(s, x2, y2, x2-11*math.cos(ang+offset), y2-11*math.sin(ang+offset), color, 2)

# 1. Cover
s = slide('Cover', 'AI research. Evidence you can check.')
chrome(s)
text(s, 'AI research.\nEvidence you\ncan check.', 64, 145, 745, 88, 'serif', h=330, leading=1.04)
text(s, 'For students and professionals.', 66, 510, 740, 29)
text(s, 'Verify the claims in AI-assisted work\nagainst the academic evidence.', 66, 561, 735, 24, color=MUTED, leading=1.32)
image(s, ASSETS / 'proof-logo.png', 835, 139, 318, 318)
rect(s, 845, 500, 327, 91, SAGE)
text(s, 'Claim / paper / passage', 868, 531, 284, 23, 'medium', FOREST)
s['notes'] = "Template slide 1, cover. Suggested timing: 15 seconds.\nProof helps students and professionals check the evidence behind AI-assisted research. Our focus is whether the academic source supports the sentence, with passages users can inspect. This is the first pitch draft, not a claim of exhaustive verification or commercial traction."

# 2. Opening scenario, visibly illustrative
s = slide('Opening scenario', 'The answer sounded right.')
chrome(s)
text(s, 'The answer sounded right.', 64, 113, 1120, 65, 'serif')
rect(s, 64, 218, 557, 291, WHITE, LINE)
rect(s, 659, 218, 557, 291, SAGE)
text(s, 'An AI-assisted draft', 92, 245, 500, 20, 'medium', MUTED)
text(s, '"The intervention\ndoubled productivity."', 92, 296, 500, 36, 'serif', leading=1.2)
text(s, '2x', 92, 409, 275, 58, 'serif', AMBER)
text(s, 'The academic passage', 688, 245, 499, 20, 'medium', FOREST)
text(s, '"Tasks were completed\n18% faster, on average."', 688, 296, 499, 36, 'serif', leading=1.2)
text(s, '18%', 688, 409, 270, 58, 'serif', FOREST)
text(s, 'An assignment. A client report.\nThe same unchecked claim.', 65, 546, 1060, 33, 'serif', leading=1.25)
text(s, 'Illustrative scenario. The quoted study and draft are fictional.', 64, 640, 1152, 14, color=MUTED)
s['notes'] = "Template slide 2, opening story or example. Suggested timing: 30 seconds.\nImagine asking an AI assistant to explain a study. The answer is fluent and includes a citation, so it looks usable. A student might put it into an essay; an analyst might use it in a client report. But the original passage says 18% faster on a measured task, not twice the productivity. This is a fictional illustration, not my personal experience or an actual study. It introduces the gap between a convincing answer and evidence that supports it."

# 3. Problem and audience
s = slide('Problem and audience', 'AI output can look credible. The evidence may not be.')
chrome(s)
text(s, 'AI output can look credible.\nThe evidence may not be.', 64, 113, 1152, 64, 'serif', leading=1.05)
line(s, 640, 291, 640, 470)
text(s, 'Academic access is incomplete.', 64, 298, 541, 28, 'medium')
text(s, 'Relevant papers can be missed.\nFull text may be unavailable.', 64, 357, 541, 27, color=MUTED, leading=1.35)
text(s, 'Hallucinations survive formatting.', 687, 298, 529, 28, 'medium')
text(s, 'References can be invented.\nReal papers can be misrepresented.', 687, 357, 529, 27, color=MUTED, leading=1.35)
rect(s, 64, 511, 1152, 116, SAGE)
for x, title, sub in [(89, 'Students', 'Essays and theses'), (375, 'Researchers', 'Manuscripts'), (662, 'Consultants', 'Client reports'), (948, 'Analysts', 'Evidence briefs')]:
    text(s, title, x, 535, 251, 25, 'medium', FOREST)
    text(s, sub, x, 578, 251, 19, color=MUTED)
text(s, 'Paper search exists. Claims still need checking against the original evidence.', 64, 640, 1152, 14, color=MUTED)
s['notes'] = "Template slide 3, problem statement. Suggested timing: 30 seconds.\nAI-assisted research can contain invented sources, inaccurate reference details, or unsupported interpretations of real papers. Search-enabled tools already find academic work, so the claim is not that AI cannot search. The problem is that discovery, full-text access, and interpretation can fail. Our proposed users include students, researchers, consultants, and analysts whose writing depends on academic evidence. These are target segments, not validated paying customers.\nSource: OpenAI deep-research launch documentation describes browsing alongside hallucination and inference limitations. https://openai.com/index/introducing-deep-research/"

# 4. One study with exact counts and visible scope
s = slide('Research evidence', '99 of 176 citations were fabricated or inaccurate in one study.')
chrome(s)
text(s, 'The citation looked real.', 64, 113, 1152, 64, 'serif')
text(s, '56%', 64, 249, 419, 122, 'serif', FOREST, h=145)
text(s, 'fabricated or inaccurate\nin one controlled study', 68, 422, 417, 27, leading=1.3)
text(s, '176 references generated by GPT-4o', 533, 249, 660, 25, 'medium')
counts = [(35, AMBER, WHITE), (64, 'B6CDBE', INK), (77, FOREST, WHITE)]
x = 533
for count, color, tc in counts:
    width = 660*count/176
    rect(s, x, 307, width, 84, color, radius=0)
    text(s, str(count), x+5, 329, width-10, 34, 'medium', tc, align='center')
    x += width
for yy, color, label in [(420, AMBER, '35 fabricated references'), (467, 'B6CDBE', '64 real references with errors'), (514, FOREST, '77 accurate references')]:
    rect(s, 535, yy+5, 18, 18, color, radius=3)
    text(s, label, 573, yy, 620, 24)
text(s, 'June 2025 · Six mental-health literature reviews · Straightforward prompts', 64, 595, 1152, 18, color=MUTED)
text(s, '99 / 176 = 56.3%. Study-specific result, not an error rate for every AI tool.', 64, 628, 935, 14, color=MUTED)
text(s, 'Linardon et al., 2025', 1000, 632, 216, 13, color=FOREST, align='right', url='https://doi.org/10.2196/80371')
s['notes'] = "Template slide 4, statistics and research. Suggested timing: 35 seconds.\nA June 2025 experiment prompted GPT-4o to generate six literature reviews on mental-health topics. The authors verified 176 references. They reported 35 fabricated references and 64 real references with bibliographic errors. The remaining 77 were accurate. We calculate 99 divided by 176, or 56.3%, and round it to 56% on the slide. These are citation errors, not a measured rate of unsupported semantic claims. This is a small, model-specific experiment under straightforward prompting, not a benchmark of all current AI tools. The study validates a failure mode; it does not establish demand or Proof's accuracy.\nSource: Linardon et al., JMIR Mental Health, 2025, DOI 10.2196/80371. https://mental.jmir.org/2025/1/e80371"

# 5. Template's question transition
s = slide('Question', 'How can we check AI-assisted research before we rely on it?', FOREST)
chrome(s)
text(s, 'How can we check\nAI-assisted research\nbefore we rely on it?', 98, 186, 1084, 81, 'serif', WHITE, leading=1.09, align='center', h=287)
line(s, 474, 521, 806, 521, LIGHT, 3)
text(s, 'Before it reaches a reader, reviewer, or client.', 121, 568, 1038, 26, color=LIGHT, align='center')
s['notes'] = "Template slide 5, turn the problem into a question. Suggested timing: 10 seconds.\nHow can we check AI-assisted research before we rely on it? The question moves us from the risk to a concrete review workflow. Pause briefly, then reveal the product."

# 6. Current UI plus a clearly illustrative review result
s = slide('Solution · Demo follows', 'Proof compares the claim with the academic evidence.')
chrome(s)
text(s, 'Proof compares the claim\nwith the academic evidence.', 64, 113, 1152, 60, 'serif', leading=1.07)
steps = [('Bring your draft.', 'Paste text or upload a document.'), ('Find or add papers.', 'Locate eligible academic sources.'), ('Inspect each finding.', 'Read the evidence before revising.')]
for i, (title, sub) in enumerate(steps):
    yy = 295 + 93*i
    rect(s, 64, yy+1, 38, 38, SAGE, radius=19)
    text(s, str(i+1), 64, yy+6, 38, 22, 'medium', FOREST, align='center')
    text(s, title, 119, yy, 377, 27, 'medium')
    text(s, sub, 119, yy+42, 377, 19, color=MUTED)
text(s, 'Current interface', 536, 270, 672, 16, color=MUTED)
image(s, ASSETS / 'proof-interface.jpg', 534, 307, 682, 162.508)
rect(s, 534, 490, 682, 123, WHITE, LINE)
text(s, 'Illustrative finding', 559, 510, 365, 16, color=MUTED)
rect(s, 1035, 506, 157, 29, AMBER_BG, radius=14)
text(s, 'Overstated', 1042, 511, 143, 15, 'medium', AMBER, align='center')
text(s, 'Draft says 2x. Passage says 18% faster.', 559, 552, 632, 27, 'serif')
text(s, 'Live demo next: find a paper, check a claim, inspect the passage.', 64, 638, 1152, 17, color=FOREST)
s['notes'] = "Template slide 6, showcase the solution. Suggested timing: 25 seconds, followed by a 60-second live demo.\nProof accepts a draft, helps find or add academic papers, and brings evidence alongside findings. Show that users can inspect the original passage and that incomplete evidence stays unverified. The screenshot is an existing local interface reference. The review card below it is illustrative, not a screenshot of a completed audit.\nLIVE DEMO, as requested by the template after slide 6:\n1. Open the current Proof app and show Check writing, Find sources, and My sources.\n2. Use a prepared, publicly accessible research article and a draft with one faithful claim plus one planted overstatement. Do not use private student or client material.\n3. Run a check, open a finding, and read the passage with the audience. Show the source-access label and an unverified case when available.\n4. Explain that the user decides what to change. Do not promise a correct live result until the demo has been rehearsed.\nCurrent implementation includes source identity and policy checks with limited journal coverage. Full-text retrieval, extraction, and AI judgments can still fail.\nImplementation references: README.md, src/Classroom.tsx, server/discovery.ts, server/scholarly.ts, server/judge.ts. Local demo: http://127.0.0.1:4317/app"

# 7. Real stack from local source
s = slide('Tech stack', 'A workflow built around the source evidence.')
chrome(s)
text(s, 'A workflow built around\nthe source evidence.', 64, 113, 1152, 61, 'serif', leading=1.07)
for xx, title, sub in [(64, 'React + TypeScript', 'Vite\nBrowser draft storage'), (488, 'Node.js + Express', 'PDF / DOCX parsing\nReview jobs and sessions'), (912, 'Jev / TypeSafe', 'Typed judgments\nClaims against passages')]:
    rect(s, xx, 281, 304, 159, WHITE, LINE)
    text(s, title, xx+24, 306, 256, 25, 'medium')
    text(s, sub, xx+24, 357, 256, 21, color=MUTED, leading=1.3)
arrow(s, 384, 361, 472, 361)
arrow(s, 808, 361, 896, 361)
rect(s, 488, 480, 728, 86, SAGE)
text(s, 'Academic data and source checks', 512, 493, 680, 17, 'medium', FOREST)
text(s, 'OpenAlex · Crossref · DOAJ · Europe PMC', 512, 526, 680, 23)
arrow(s, 641, 477, 641, 442)
rect(s, 64, 598, 1152, 42, FOREST, radius=10)
text(s, 'Missing or conflicting evidence stays unverified.', 83, 607, 1114, 20, color=WHITE)
s['notes'] = "Template slide 7, tech stack. Suggested timing: 35 seconds.\nThe browser uses React, TypeScript, and Vite. The Node.js and Express server parses uploaded documents and runs review jobs. OpenAlex supplies discovery candidates, Crossref resolves article identity and metadata, DOAJ helps confirm journal peer-review policy, and Europe PMC supplies available full-text evidence. Jev through TypeSafe answers typed questions about claims and evidence. Code checks the responses and preserves uncertainty when retrieval, extraction, or judgments fail. Current public-source checks deliberately exclude articles outside the verifiable policy and full-text set. A failed lookup does not prove a source is fabricated. No provider credentials or private inputs are included in the deck.\nImplementation references: package.json, server/discovery.ts, server/sources.ts, server/scholarly.ts, server/judge.ts, server/evidence.ts, src/class-store.ts."

# 8. Actual visual system plus a closing sentence
s = slide('Design system', 'Trust starts with what you can inspect.')
chrome(s)
text(s, 'Trust starts with\nwhat you can inspect.', 64, 113, 648, 63, 'serif', leading=1.06)
text(s, 'Aa', 64, 313, 218, 105, 'serif', h=125)
text(s, 'Aa', 328, 313, 219, 105, 'sans', h=125)
text(s, 'Newsreader', 68, 452, 224, 24, 'medium')
text(s, 'Headings and quotations', 68, 491, 251, 18, color=MUTED)
text(s, 'DM Sans', 332, 452, 224, 24, 'medium')
text(s, 'Controls and body text', 332, 491, 251, 18, color=MUTED)
for xx, color, label in [(723, INK, 'Ink'), (851, FOREST, 'Forest'), (979, SAGE, 'Sage'), (1107, PAPER, 'Paper')]:
    rect(s, xx, 157, 109, 102, color, LINE if color==PAPER else None, radius=12)
    text(s, label, xx, 280, 109, 19, 'medium')
    text(s, '#'+color.lower(), xx, 314, 109, 14, color=MUTED)
rect(s, 723, 389, 493, 171, WHITE, LINE)
text(s, 'Evidence beside the claim.\nUncertainty stays visible.', 747, 410, 445, 26, 'serif', leading=1.25)
rect(s, 747, 498, 202, 36, SAGE, radius=8)
text(s, 'Source passage', 755, 506, 186, 17, 'medium', FOREST, align='center')
rect(s, 971, 498, 220, 36, 'F3F7F4', LINE, radius=8)
text(s, 'Unverified', 979, 506, 204, 17, color=MUTED, align='center')
text(s, 'Proof. Check the research before you rely on it.', 64, 601, 1152, 30, 'serif', FOREST)
s['notes'] = "Template slide 8, design system and closing. Suggested timing: 25 seconds.\nProof's current visual system uses Newsreader headings, DM Sans body text, forest green, sage, and paper backgrounds. Evidence passages and uncertainty labels are part of the review experience. Close with the product promise: check the research before you rely on it. Invite students and professional research teams to test the workflow. No customer count, revenue, accuracy guarantee, or paid-pilot claim has been supplied or added.\nDesign references: src/palette.css, src/landing.css, src/styles.css, public/design-system/app-reference.jpg."

# Shared layout validation uses the same font metrics as the PDF.
checks = []
font_chars = {name: set(TTFont(ASSETS / f'{family}-latin-{weight}-{style}.ttf').getBestCmap()) for name, (family, weight, style, display) in FONTS.items()}
for number, s in enumerate(slides, 1):
    for e in s['elements']:
        if e['type'] == 'text':
            missing = {ch for ch in e['text'] if ch != '\n' and ord(ch) not in font_chars[e['font']]}
            if missing:
                checks.append(f'Slide {number}: missing glyphs in {e["font"]}: {missing}')
            for value in e['text'].split('\n'):
                width = pdfmetrics.stringWidth(value, e['font'], e['size'])
                if width > e['w'] + 1:
                    checks.append(f'Slide {number}: text too wide ({width:.1f} > {e["w"]}): {value}')
            if e['y'] + e['h'] > H:
                checks.append(f'Slide {number}: text extends below canvas: {e["text"]}')
        if e['type'] in ('image', 'rect'):
            if e['x'] < 0 or e['y'] < 0 or e['x'] + e['w'] > W+.1 or e['y']+e['h']>H+.1:
                checks.append(f'Slide {number}: out-of-bounds {e["type"]}')
if checks:
    raise ValueError('\n'.join(checks))

# PDF: actual embedded Newsreader and DM Sans fonts, searchable text.
pdf = canvas.Canvas(str(OUT / 'proof-pitch-draft.pdf'), pagesize=(W*.75,H*.75))
pdf.setTitle('Proof | Hackathon pitch | First draft')
pdf.setAuthor('Proof')
for s in slides:
    pdf.setFillColor('#'+s['bg'])
    pdf.rect(0,0,W*.75,H*.75,fill=1,stroke=0)
    for e in s['elements']:
        if e['type']=='rect':
            pdf.setFillColor('#'+e['fill'])
            pdf.setStrokeColor('#'+(e['stroke'] or e['fill']))
            pdf.setLineWidth(.75)
            pdf.roundRect(e['x']*.75,(H-e['y']-e['h'])*.75,e['w']*.75,e['h']*.75,e['radius']*.75,fill=1,stroke=int(bool(e['stroke'])))
        elif e['type']=='line':
            pdf.setStrokeColor('#'+e['color']);pdf.setLineWidth(e['width']*.75)
            pdf.line(e['x1']*.75,(H-e['y1'])*.75,e['x2']*.75,(H-e['y2'])*.75)
        elif e['type']=='image':
            pdf.drawImage(e['path'], e['x']*.75,(H-e['y']-e['h'])*.75,e['w']*.75,e['h']*.75,mask='auto')
        elif e['type']=='text':
            pdf.setFont(e['font'],e['size']*.75);pdf.setFillColor('#'+e['color'])
            for i, value in enumerate(e['text'].split('\n')):
                yy = (H-e['y']-e['size']*.82-i*e['size']*e['leading'])*.75
                if e['align']=='right': pdf.drawRightString((e['x']+e['w'])*.75,yy,value)
                elif e['align']=='center': pdf.drawCentredString((e['x']+e['w']/2)*.75,yy,value)
                else: pdf.drawString(e['x']*.75,yy,value)
            if e.get('url'):
                pdf.linkURL(e['url'],(e['x']*.75,(H-e['y']-e['h'])*.75,(e['x']+e['w'])*.75,(H-e['y'])*.75),relative=0)
    pdf.showPage()
pdf.save()

# PPTX: native text, shapes, chart segments, connectors; two image assets.
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W/96), Inches(H/96)
prs.core_properties.title = 'Proof | Hackathon pitch | First draft'
prs.core_properties.subject = 'AI-assisted research verification for students and professionals'
prs.core_properties.author = 'Proof'
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
for s in slides:
    ps = prs.slides.add_slide(prs.slide_layouts[6])
    ps.background.fill.solid(); ps.background.fill.fore_color.rgb = RGBColor.from_string(s['bg'])
    for e in s['elements']:
        if e['type']=='rect':
            shape=ps.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e['radius'] else MSO_SHAPE.RECTANGLE, Inches(e['x']/96), Inches(e['y']/96), Inches(e['w']/96), Inches(e['h']/96))
            if e['radius']:
                shape.adjustments[0]=min(.5,e['radius']/min(e['w'],e['h']))
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(e['fill'])
            if e['stroke']:
                shape.line.color.rgb=RGBColor.from_string(e['stroke']);shape.line.width=Pt(.75)
            else: shape.line.fill.background()
        elif e['type']=='line':
            shape=ps.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e['x1']/96), Inches(e['y1']/96), Inches(e['x2']/96), Inches(e['y2']/96))
            shape.line.color.rgb=RGBColor.from_string(e['color']);shape.line.width=Pt(e['width']*.75)
        elif e['type']=='image':
            ps.shapes.add_picture(e['path'], Inches(e['x']/96), Inches(e['y']/96), Inches(e['w']/96), Inches(e['h']/96))
        elif e['type']=='text':
            shape=ps.shapes.add_textbox(Inches(e['x']/96), Inches(e['y']/96), Inches(e['w']/96), Inches(e['h']/96))
            tf=shape.text_frame;tf.clear();tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.vertical_anchor=MSO_ANCHOR.TOP
            for i, value in enumerate(e['text'].split('\n')):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
                p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[e['align']]
                p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=e['leading']
                run=p.add_run();run.text=value
                run.font.name=FONTS[e['font']][3]
                run.font.size=Pt(e['size']*.75)
                run.font.italic=e['font']=='italic';run.font.bold=False
                run.font.color.rgb=RGBColor.from_string(e['color'])
                if e.get('url'): run.hyperlink.address=e['url']
        # Theme shape styles can add shadows that are absent from the source layout.
        for style in list(shape._element.findall(qn('p:style'))):
            shape._element.remove(style)
        sp_pr = shape._element.find(qn('p:spPr'))
        if sp_pr is not None and sp_pr.find(qn('a:effectLst')) is None:
            sp_pr.append(OxmlElement('a:effectLst'))
    ps.notes_slide.notes_text_frame.text=s['notes']
prs.save(OUT / 'proof-pitch-draft.pptx')

# HTML: static SVG lettering also survives Postplan's font-src policy.
# Reuse each glyph outline, so font preservation fits its HTML size limit.
outline_fonts = {name: TTFont(ASSETS / f'{family}-latin-{weight}-{style}.ttf') for name, (family, weight, style, display) in FONTS.items()}
font_ids = {'serif':'r','italic':'i','sans':'s','medium':'m'}
glyph_defs = []
for name, ff in outline_fonts.items():
    characters = sorted({ch for s in slides for e in s['elements'] if e['type']=='text' and e['font']==name for ch in e['text'] if ch!='\n'})
    cmap = ff.getBestCmap()
    glyphset = ff.getGlyphSet()
    for ch in characters:
        pen = SVGPathPen(glyphset)
        glyphset[cmap[ord(ch)]].draw(pen)
        commands = pen.getCommands()
        if commands:
            glyph_defs.append(f'<path id="{font_ids[name]}{ord(ch)}" d="{commands}"/>')
css_fonts=[]
parts=['<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="color-scheme" content="light"><title>Proof | First pitch deck draft</title><style>'+''.join(css_fonts)+'''
*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#e7efe9;color:#1d3027;font:16px/1.6 "sans",Arial,sans-serif}a{color:inherit;text-underline-offset:4px}nav{position:sticky;top:0;z-index:10;display:flex;justify-content:space-between;align-items:center;gap:20px;padding:14px 32px;background:#fafbf9eF;border-bottom:1px solid #dce5df;backdrop-filter:blur(8px)}nav strong{font:28px "serif",Georgia,serif}nav .links{display:flex;gap:16px;align-items:center}nav a{text-decoration:none;padding:2px 6px}a:focus-visible,summary:focus-visible{outline:3px solid #8fc5ab;outline-offset:4px}.deck{max-width:1400px;margin:32px auto;padding:0 32px}.frame{scroll-margin-top:100px;margin-bottom:32px;background:#fafbf9;border-radius:12px;overflow:hidden;box-shadow:0 10px 32px #1d30270c}.frame svg{display:block;width:100%;height:auto}.frame details{padding:15px 24px;border-top:1px solid #dce5df;font-size:14px}.frame summary{cursor:pointer}.notes{max-width:1000px;white-space:pre-wrap;padding-top:12px;color:#5c6d63;font:15px/1.65 "sans",Arial,sans-serif}.sources{padding:24px;background:#fafbf9;border-radius:12px}.sources h2{font:30px "serif",Georgia,serif;margin:0 0 12px}.sources p{margin:8px 0;font-size:14px}.source-list{display:flex;gap:20px;flex-wrap:wrap;margin-top:18px}.source-list a{font-size:14px}.end{padding:24px 0 12px;color:#5c6d63;font-size:13px}@media(max-width:640px){nav{padding:12px 16px;gap:8px;flex-wrap:wrap}nav .links{gap:6px;font-size:14px}.deck{padding:0 12px;margin-top:16px}.frame{margin-bottom:20px;border-radius:8px}.frame details{padding:12px 16px}.sources{padding:18px}}@media(prefers-reduced-motion:reduce){html{scroll-behavior:auto}}@media print{@page{size:13.333in 7.5in;margin:0}nav,.sources,.end,details{display:none}.deck{padding:0;margin:0;max-width:none}.frame{border-radius:0;margin:0;box-shadow:none;break-after:page}.frame:last-of-type{break-after:auto}.frame svg{width:100%;height:100vh}}
</style></head><body><nav aria-label="Slide navigation"><strong>proof.</strong><div class="links"><span>First pitch draft</span>''']
for i in range(1,9): parts.append(f'<a href="#slide-{i}" aria-label="Go to slide {i}">{i}</a>')
parts.append('<a href="#sources">Sources</a></div></nav><svg xmlns="http://www.w3.org/2000/svg" width="0" height="0" aria-hidden="true" style="position:absolute"><defs>'+''.join(glyph_defs)+'</defs></svg><main class="deck">')
for number, s in enumerate(slides,1):
    body_text=' '.join(e['text'] for e in s['elements'] if e['type']=='text')
    parts.append(f'<section class="frame" id="slide-{number}" aria-label="Slide {number}: {escape(s["section"])}"><svg viewBox="0 0 1280 720" xmlns="http://www.w3.org/2000/svg" role="img" aria-labelledby="title-{number} desc-{number}"><title id="title-{number}">{escape(s["title"])}</title><desc id="desc-{number}">{escape(body_text)}</desc><rect width="1280" height="720" fill="#{s["bg"]}"/>')
    for e in s['elements']:
        if e['type']=='rect':
            parts.append(f'<rect x="{e["x"]}" y="{e["y"]}" width="{e["w"]}" height="{e["h"]}" rx="{e["radius"]}" fill="#{e["fill"]}" stroke="#{e["stroke"] or e["fill"]}" stroke-width="{1 if e["stroke"] else 0}"/>')
        elif e['type']=='line':
            parts.append(f'<line x1="{e["x1"]}" y1="{e["y1"]}" x2="{e["x2"]}" y2="{e["y2"]}" stroke="#{e["color"]}" stroke-width="{e["width"]}"/>')
        elif e['type']=='image':
            mime='image/png' if e['path'].endswith('.png') else 'image/jpeg'
            data=base64.b64encode(Path(e['path']).read_bytes()).decode()
            parts.append(f'<image href="data:{mime};base64,{data}" x="{e["x"]}" y="{e["y"]}" width="{e["w"]}" height="{e["h"]}"/>')
        elif e['type']=='text':
            if e.get('url'):parts.append(f'<a href="{escape(e["url"],quote=True)}" aria-label="Research source">')
            for i, value in enumerate(e['text'].split('\n')):
                yy=e['y']+e['size']*.82+i*e['size']*e['leading']
                width=pdfmetrics.stringWidth(value,e['font'],e['size'])
                xx=e['x'] if e['align']=='left' else e['x']+(e['w']-width)/2 if e['align']=='center' else e['x']+e['w']-width
                ff=outline_fonts[e['font']]; cmap=ff.getBestCmap()
                scale=e['size']/ff['head'].unitsPerEm
                parts.append(f'<g fill="#{e["color"]}" transform="translate({xx:.3f} {yy:.3f}) scale({scale:.6f} {-scale:.6f})" aria-label="{escape(value,quote=True)}">')
                advance=0
                for ch in value:
                    if ch!=' ':
                        parts.append(f'<use href="#{font_ids[e["font"]]}{ord(ch)}" x="{advance}"/>')
                    advance+=ff['hmtx'][cmap[ord(ch)]][0]
                parts.append('</g>')
            if e.get('url'):parts.append('</a>')
    parts.append('</svg><details><summary>Speaker notes and template mapping</summary><div class="notes">'+escape(s['notes'])+'</div></details></section>')
parts.append('''<section class="sources" id="sources"><h2>Sources and draft scope</h2><p>Eight slides follow the supplied Hackathon Pitch Template. The live-demo walkthrough is in slide six's speaker notes. Suggested speaking time is about 3 minutes 25 seconds, plus a 60-second demo.</p><p>Public, redacted draft. No private student or client content. Market segments are proposed; no traction, revenue, pricing, or accuracy guarantee is claimed. Product images are a local interface reference and clearly labeled illustrations.</p><p>Research chart: 35 fabricated + 64 real with errors + 77 accurate = 176 references. The combined 56% figure is calculated from the reported counts. It is specific to GPT-4o, six mental-health literature reviews, and straightforward prompts in June 2025.</p><div class="source-list"><a href="https://doi.org/10.2196/80371">Linardon et al., JMIR Mental Health, 2025</a><a href="https://openai.com/index/introducing-deep-research/">OpenAI deep-research documentation</a><a href="https://elicit.com/">Elicit</a><a href="https://scite.ai/getting-started">Scite</a></div></section><p class="end">Proof · First pitch deck draft · 30 September 2026 · Editable PowerPoint and PDF supplied separately.</p></main></body></html>''')
(OUT / 'proof-pitch-draft.html').write_text(''.join(parts))
(OUT / 'slide-content.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2))
(OUT / 'speaker-notes.md').write_text('# Proof pitch speaker notes\n\n'+ '\n\n'.join(f'## {i}. {s["section"]}\n\n{s["notes"]}' for i,s in enumerate(slides,1)))
(OUT / 'README.md').write_text('''# Proof pitch, first draft

Eight slides follow the supplied Hackathon Pitch Template. The live demo follows slide six and has a walkthrough in its notes.

- `proof-pitch-draft.pptx`: editable PowerPoint, with speaker notes and native text, diagrams, and chart segments.
- `proof-pitch-draft.pdf`: searchable PDF with embedded fonts.
- `proof-pitch-draft.html`: self-contained browser deck with navigation, collapsible notes, and sources. Slide lettering uses SVG outlines to preserve the fonts under the host's security policy; accessible slide descriptions retain the text.
- `speaker-notes.md`: speaking script and source context.
- `assets/`: Newsreader and DM Sans font files and the two image assets. Install the fonts locally if your presentation app substitutes them. The PDF and HTML include their fonts.
- `build_deck.py`: local rebuild script using Python, fontTools, Pillow, ReportLab, and python-pptx. It reads the project's current logo, interface reference, and Fontsource files.

Rebuild from this directory with `python3 build_deck.py`.

This is a public, redacted draft. All audience segments are proposed. No commercial traction, prices, customer counts, or general accuracy rates are invented. The example on slide two and review card on slide six are illustrative. The study chart is based on reported counts, not the paper's imprecise headline wording. The PDF template itself is not republished.
''')
print(json.dumps({'slides':len(slides),'pptx':str(OUT/'proof-pitch-draft.pptx'),'pdf':str(OUT/'proof-pitch-draft.pdf'),'html_bytes':(OUT/'proof-pitch-draft.html').stat().st_size,'layout_checks':'pass'},indent=2))
